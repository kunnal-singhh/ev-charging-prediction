"""
Build the Progress Presentation for EVConnect Minor Project as per supervisor's interim review guidelines.
Strictly complies with:
- NIT Raipur PPT Template & Slide Limits
- Progress / Interim Presentation narrative (Model 1: Logistic Regression -> Model 2: Random Forest -> Ongoing Exploration)
- EXACT metrics from data/processed/model_comparison_v1.csv
- ZERO mention of XGBoost anywhere in slides, notes, or metadata
- Output: EVConnect_Minor_Project_Progress_Presentation_NIT_Raipur.pptx (Desktop & Workspace)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

TEMPLATE_PATH = r"C:\Users\kunal singh\OneDrive\Desktop\MCA-V-Sem- Project Presentation Template.pptx"
DESKTOP_OUT = r"C:\Users\kunal singh\OneDrive\Desktop\EVConnect_Minor_Project_Progress_Presentation_NIT_Raipur.pptx"
WORKSPACE_OUT = r"EVConnect_Minor_Project_Progress_Presentation_NIT_Raipur.pptx"

# Color Palette constants (Academic, matching NIT Raipur theme)
COLOR_NAVY = RGBColor(0, 32, 96)       # #002060 Department Navy
COLOR_DARK = RGBColor(15, 23, 42)      # #0f172a Deep slate text
COLOR_MUTED = RGBColor(71, 85, 105)    # #475569 Secondary text
COLOR_BLUE = RGBColor(2, 132, 199)     # #0284c7 Accent blue
COLOR_GREEN = RGBColor(16, 185, 129)   # #10b981 Success green

def configure_title_slide(slide):
    """Slide 1: Title Slide"""
    title = slide.shapes.title
    title.text = "EVConnect: Machine Learning-Based Prediction of EV Charging Station Occupancy"
    title.text_frame.paragraphs[0].font.size = Pt(36)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = COLOR_NAVY
    
    # Subtitle placeholder
    sub = slide.placeholders[1]
    tf = sub.text_frame
    tf.word_wrap = True
    tf.clear()
    
    p0 = tf.paragraphs[0]
    p0.text = "Minor Project Progress Presentation"
    p0.font.size = Pt(20)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_BLUE
    p0.space_after = Pt(10)
    
    p1 = tf.add_paragraph()
    p1.text = "Presented by\nKunal Singh"
    p1.font.size = Pt(24)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_NAVY
    p1.space_after = Pt(6)
    
    p2 = tf.add_paragraph()
    p2.text = "Roll No: [Insert Roll No / 23223xxx]\nMCA - V Semester\n(Session: 2024-2026)"
    p2.font.size = Pt(16)
    p2.font.color.rgb = COLOR_DARK
    
    # Supervisor box
    for s in slide.shapes:
        if s.name == "TextBox 5":
            stf = s.text_frame
            stf.clear()
            sp0 = stf.paragraphs[0]
            sp0.text = "Guided by:"
            sp0.font.size = Pt(15)
            sp0.font.color.rgb = COLOR_MUTED
            
            sp1 = stf.add_paragraph()
            sp1.text = "[Project Supervisor Name]"
            sp1.font.size = Pt(19)
            sp1.font.bold = True
            sp1.font.color.rgb = COLOR_NAVY
            
            sp2 = stf.add_paragraph()
            sp2.text = "Supervisor\nDepartment of Computer Applications\nNational Institute of Technology Raipur"
            sp2.font.size = Pt(14)
            sp2.font.color.rgb = COLOR_DARK

    slide.notes_slide.notes_text_frame.text = (
        "Good morning respected members of the evaluation jury and my project supervisor. "
        "Today, I present my Minor Project Progress Presentation on EVConnect: Machine Learning-Based "
        "Prediction of EV Charging Station Occupancy. In this presentation, I will demonstrate the work "
        "completed so far: data preparation, problem formulation, feature engineering, and our comparative "
        "investigation of our first two machine learning models—Logistic Regression and Random Forest."
    )

def configure_guideline_slide(slide):
    """Slide 2: Sequence of slides to be followed (Retain from template)"""
    title = slide.shapes.title
    title.text = "Sequence of slides to be followed"
    title.text_frame.paragraphs[0].font.color.rgb = COLOR_NAVY
    title.text_frame.paragraphs[0].font.bold = True
    
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.word_wrap = True
    tf.clear()
    
    sequence_items = [
        "1.  Declaration",
        "2.  Problem Formulation and Definition (one slide)",
        "3.  Motivation (one slide)",
        "4.  Contribution (one slide)",
        "5.  Methodology/ Proposed Algorithm (maximum five slides)",
        "6.  Design/ Architecture/ Workflow (maximum five slides)",
        "7.  Screen Shots/Testing & Implementation/Results and Discussion (maximum five slides)",
        "8.  Conclusion and Future scope (one slide)",
        "9.  References/ Bibliography (one slide)"
    ]
    
    for i, item in enumerate(sequence_items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(18)
        p.font.color.rgb = COLOR_DARK
        p.space_after = Pt(10)
        
    slide.notes_slide.notes_text_frame.text = (
        "This progress review strictly follows the departmental sequence specified by the Department of "
        "Computer Applications at NIT Raipur, maintaining single-slide and maximum five-slide boundaries."
    )

def configure_contents_slide(slide):
    """Slide 3: Contents"""
    title = slide.shapes.title
    title.text = "Contents:"
    title.text_frame.paragraphs[0].font.color.rgb = COLOR_NAVY
    title.text_frame.paragraphs[0].font.bold = True
    
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.word_wrap = True
    tf.clear()
    
    items = [
        "1.  Declaration",
        "2.  Problem Formulation and Definition",
        "3.  Motivation",
        "4.  Work Completed So Far (Contributions)",
        "5.  Methodology: Data Preprocessing & Feature Engineering",
        "6.  Design & Experimental Research Workflow",
        "7.  Model 1 (Logistic Regression) & Model 2 (Random Forest) Results",
        "8.  Current Observations & Future Research Scope",
        "9.  References / Bibliography"
    ]
    
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(18)
        p.font.color.rgb = COLOR_DARK
        p.space_after = Pt(10)
        
    slide.notes_slide.notes_text_frame.text = (
        "Here is the outline of today's progress presentation: we will examine the problem definition, "
        "data preprocessing pipeline, the 8 features engineered, the reasoning for selecting Logistic Regression "
        "first, why we transitioned to Random Forest, our empirical findings, and our planned next steps."
    )

def add_two_column_slide(prs, title_text, bullets, image_path, notes_text):
    """Creates a standardized two-column slide (bullets left, image right)"""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.size = Pt(28)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = COLOR_NAVY
    
    # Left column for bullets
    body = slide.placeholders[1]
    body.left = Inches(0.85)
    body.top = Inches(1.55)
    body.width = Inches(6.35)
    body.height = Inches(5.15)
    
    tf = body.text_frame
    tf.word_wrap = True
    tf.clear()
    
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        
        if isinstance(b, tuple):
            run1 = p.add_run()
            run1.text = b[0] + ": "
            run1.font.bold = True
            run1.font.size = Pt(15.5)
            run1.font.color.rgb = COLOR_NAVY
            
            run2 = p.add_run()
            run2.text = b[1]
            run2.font.bold = False
            run2.font.size = Pt(14.5)
            run2.font.color.rgb = COLOR_DARK
        else:
            run = p.add_run()
            run.text = b
            run.font.size = Pt(15)
            run.font.color.rgb = COLOR_DARK
            
    # Right column for image
    if image_path and os.path.exists(image_path):
        slide.shapes.add_picture(image_path, Inches(7.4), Inches(1.55), width=Inches(5.15))
        
    if notes_text:
        slide.notes_slide.notes_text_frame.text = notes_text
        
    return slide

def add_full_width_slide(prs, title_text, items, notes_text):
    """Creates a full-width slide with formatted paragraphs/cards"""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.size = Pt(28)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = COLOR_NAVY
    
    body = slide.placeholders[1]
    body.left = Inches(0.85)
    body.top = Inches(1.55)
    body.width = Inches(11.5)
    body.height = Inches(5.15)
    
    tf = body.text_frame
    tf.word_wrap = True
    tf.clear()
    
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        
        if isinstance(item, tuple):
            run1 = p.add_run()
            run1.text = item[0] + (": " if len(item[0]) < 35 else "\n")
            run1.font.bold = True
            run1.font.size = Pt(16.5)
            run1.font.color.rgb = COLOR_NAVY
            
            run2 = p.add_run()
            run2.text = item[1]
            run2.font.bold = False
            run2.font.size = Pt(15)
            run2.font.color.rgb = COLOR_DARK
        else:
            run = p.add_run()
            run.text = item
            run.font.size = Pt(15.5)
            run.font.color.rgb = COLOR_DARK
            
    if notes_text:
        slide.notes_slide.notes_text_frame.text = notes_text
        
    return slide

def add_two_column_cards_slide(prs, title_text, col1_title, col1_items, col2_title, col2_items, notes_text):
    """Creates side-by-side structured card comparison slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.size = Pt(28)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = COLOR_NAVY
    
    body1 = slide.placeholders[1]
    body1.left = Inches(0.85)
    body1.top = Inches(1.55)
    body1.width = Inches(5.6)
    body1.height = Inches(5.15)
    
    tf1 = body1.text_frame
    tf1.word_wrap = True
    tf1.clear()
    
    h1 = tf1.paragraphs[0]
    h1.text = col1_title
    h1.font.bold = True
    h1.font.size = Pt(19)
    h1.font.color.rgb = COLOR_NAVY
    h1.space_after = Pt(10)
    
    for item in col1_items:
        p = tf1.add_paragraph()
        p.space_after = Pt(10)
        if isinstance(item, tuple):
            r1 = p.add_run()
            r1.text = item[0] + ": "
            r1.font.bold = True
            r1.font.size = Pt(15)
            r1.font.color.rgb = COLOR_NAVY
            r2 = p.add_run()
            r2.text = item[1]
            r2.font.size = Pt(14)
            r2.font.color.rgb = COLOR_DARK
        else:
            r = p.add_run()
            r.text = item
            r.font.size = Pt(14.5)
            r.font.color.rgb = COLOR_DARK
            
    tx_box = slide.shapes.add_textbox(Inches(6.8), Inches(1.55), Inches(5.6), Inches(5.15))
    tf2 = tx_box.text_frame
    tf2.word_wrap = True
    
    h2 = tf2.paragraphs[0]
    h2.text = col2_title
    h2.font.bold = True
    h2.font.size = Pt(19)
    h2.font.color.rgb = COLOR_BLUE
    h2.space_after = Pt(10)
    
    for item in col2_items:
        p = tf2.add_paragraph()
        p.space_after = Pt(10)
        if isinstance(item, tuple):
            r1 = p.add_run()
            r1.text = item[0] + ": "
            r1.font.bold = True
            r1.font.size = Pt(15)
            r1.font.color.rgb = COLOR_BLUE
            r2 = p.add_run()
            r2.text = item[1]
            r2.font.size = Pt(14)
            r2.font.color.rgb = COLOR_DARK
        else:
            r = p.add_run()
            r.text = item
            r.font.size = Pt(14.5)
            r.font.color.rgb = COLOR_DARK
            
    if notes_text:
        slide.notes_slide.notes_text_frame.text = notes_text
        
    return slide

def build_presentation():
    print("Loading official NIT Raipur presentation template...")
    prs = Presentation(TEMPLATE_PATH)
    
    # Configure Slide 1, 2, 3
    configure_title_slide(prs.slides[0])
    configure_guideline_slide(prs.slides[1])
    configure_contents_slide(prs.slides[2])
    
    # Remove template slides 4-14
    print("Clearing template placeholder slides 4 to 14...")
    while len(prs.slides) > 3:
        rId = prs.slides._sldIdLst[3].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[3]
        
    # =========================================================================
    # SLIDE 4: Declaration (Section 1 — 1 slide)
    # =========================================================================
    print("Building Slide 4: Declaration...")
    decl_paragraphs = [
        ("CANDIDATE'S DECLARATION", 
         "I hereby declare that the minor project work entitled 'EVConnect: Machine Learning-Based Prediction of EV Charging Station Occupancy' is an authentic record of my ongoing original research work carried out as part of the MCA Minor Project under the supervision and guidance of my Project Supervisor, Department of Computer Applications, National Institute of Technology (NIT), Raipur."),
        ("ACADEMIC ETHICS & PROGRESS COMMITMENT",
         "I further declare that this interim presentation reports genuine experimental progress conducted on real-world telemetry from the Adaptive Charging Network (ACN-Data). All data processing pipelines, algorithmic implementations, and comparative benchmarks are authentically performed in accordance with academic standards."),
        ("STUDENT PARTICULARS",
         "Kunal Singh  |  Roll No: [Insert Roll No / 23223xxx]\nMaster of Computer Applications (MCA) - V Semester  |  Session: 2024-2026\nDepartment of Computer Applications, National Institute of Technology Raipur")
    ]
    add_full_width_slide(
        prs, "Declaration", decl_paragraphs,
        "Good morning respected evaluation committee members. I formally declare that this presentation reflects my genuine progress in the MCA Minor Project under the Department of Computer Applications at NIT Raipur."
    )

    # =========================================================================
    # SLIDE 5: Problem Formulation and Definition (Section 2 — 1 slide)
    # =========================================================================
    print("Building Slide 5: Problem Formulation...")
    prob_bullets = [
        ("Core Prediction Goal", "Predict whether an EV charging station s will be occupied 30 minutes ahead (at time step t+1)."),
        ("Input Information at Time t", "x_{s,t} utilizes only data observed at or before timestamp t (current state, rolling 2h trend, station-hour priors, calendar)."),
        ("Prediction Target at t + 30m", "target_occupancy in {0, 1} representing the physical plug state 30 minutes into the future."),
        ("Binary Classification Semantics", "0 = Available (free charger / open spot); 1 = Occupied (EV connected and charging)."),
        ("Strict Chronological Partitioning", "Data is partitioned chronologically (Train -> Validation -> Test). Future observations are strictly excluded during training to guarantee zero data leakage.")
    ]
    add_two_column_slide(
        prs, "Problem Formulation and Definition", prob_bullets,
        "artifacts/progress_figures/fig1_problem_formulation.png",
        "Respected jury, our problem is formulated as 30-minute-ahead binary occupancy classification. The input consists strictly of information known up to time t. The target is whether the station is occupied at time t plus 30 minutes. Notice that we enforce a strict chronological boundary so that future information never leaks into training."
    )

    # =========================================================================
    # SLIDE 6: Motivation (Section 3 — 1 slide)
    # =========================================================================
    print("Building Slide 6: Motivation...")
    mot_bullets = [
        ("Dynamic EV Charging Availability", "Public EV charging availability changes rapidly and unpredictably as vehicles arrive and depart throughout the day."),
        ("Limitation of Real-Time Status", "Current charging apps show only live status. A charger showing 'available' may become occupied during the 20-30 minute drive to the station."),
        ("Value of 30-Minute Prediction", "A 30-minute lookahead forecast gives drivers actionable lead time to make informed routing and charging decisions before arriving."),
        ("Inherent Patterns in Historical Data", "Charging sessions follow distinct human routines—such as workplace commuting peaks between 09:00 and 16:00 and long dwell times."),
        ("Core Research Investigation", "Our project investigates whether machine learning algorithms can learn these spatial-temporal patterns to reliably forecast availability.")
    ]
    add_two_column_slide(
        prs, "Motivation", mot_bullets,
        "artifacts/progress_figures/fig3_class_splits.png",
        "The motivation is practical: knowing that a charger is free right now is insufficient if you have to drive 20 minutes to get there. By predicting occupancy 30 minutes ahead using historical charging telemetry, we can help drivers avoid occupied stations and reduce arrival uncertainty."
    )

    # =========================================================================
    # SLIDE 7: Contribution / Work Completed So Far (Section 4 — 1 slide)
    # =========================================================================
    print("Building Slide 7: Work Completed So Far...")
    cont_bullets = [
        ("ACN-Data Studied & Prepared", "Analyzed the open-source Adaptive Charging Network (ACN-Data) collected from 55 EVSE charging stations at Caltech."),
        ("30-Minute Prediction Grid Constructed", "Converted asynchronous raw charging session records into an aligned 30-minute temporal grid comprising 1,861,310 intervals."),
        ("Target Construction Completed", "Constructed the leak-free 30-minute forward target (target_occupancy in {0, 1}) for every station-interval pair."),
        ("Feature Engineering Completed", "Engineered 8 domain-specific features across temporal, historical lag, station prior, and session dwell categories."),
        ("Chronological Evaluation Established", "Formulated leakage-safe Train (89.34%), Validation (6.38%), and Test (4.28%) partitions."),
        ("Model 1 (Logistic Regression) Trained", "Trained balanced linear baseline model to establish a foundational reference benchmark."),
        ("Model 2 (Random Forest) Trained", "Trained non-linear tree ensemble to explore complex feature interactions and non-linear boundaries."),
        ("Comparative Analysis Performed", "Conducted rigorous empirical comparison between both models; further experimentation is currently underway.")
    ]
    add_full_width_slide(
        prs, "Contributions / Work Completed So Far", cont_bullets,
        "Here is the summary of work completed at this interim stage: we prepared the 1.86M interval dataset, constructed the 30-minute target, engineered 8 features, established a chronological evaluation setup, trained Logistic Regression, trained Random Forest, and evaluated their performance. Final model selection is reserved for ongoing work."
    )

    # =========================================================================
    # SLIDE 8: Methodology 1/3 (Data Preprocessing Pipeline)
    # =========================================================================
    print("Building Slide 8: Methodology - Data Preprocessing...")
    p1_bullets = [
        ("Raw ACN Charging Sessions", "Extracted 7,596 actual charging sessions from the Caltech facility with connection, disconnection, and station attributes."),
        ("Timestamp & Timezone Processing", "Converted UTC timestamps to America/Los_Angeles local time to properly capture local commuting and charging schedules."),
        ("30-Minute Regular Grid Alignment", "Discretized continuous sessions into synchronized 30-minute discrete slots across all 55 stations (1,861,310 total intervals)."),
        ("Target Construction (T + 30 min)", "Constructed binary label target_occupancy by evaluating physical charger connection during the subsequent 30-minute interval."),
        ("Leakage-Free Feature Processing", "Computed lag features strictly on data <= t; precalculated historical station priors exclusively from the training partition.")
    ]
    add_two_column_slide(
        prs, "Methodology: Data Preprocessing Pipeline", p1_bullets,
        "artifacts/progress_figures/fig2_preprocessing_pipeline.png",
        "Our data preprocessing pipeline transforms raw asynchronous charging sessions into a regular 30-minute occupancy grid. We converted UTC timestamps to local California time to align with human daily routines, constructed the 30-minute forward target, and enforced a strict leakage-free feature computation policy."
    )

    # =========================================================================
    # SLIDE 9: Methodology 2/3 (Feature Engineering)
    # =========================================================================
    print("Building Slide 9: Methodology - Feature Engineering...")
    f_bullets = [
        ("1. Temporal Features (4)", "local_hour (0-23), local_day_of_week (0-6), is_weekend (0/1), month (1-12). Provides diurnal commute peaks, workday rhythms, and seasonal cycles."),
        ("2. Historical Lags (2)", "previous_occupancy (occupancy at current time t), rolling_2hr_occupancy (4-step moving average). Captures charging session inertia and short-term arrival momentum."),
        ("3. Station Operational Prior (1)", "station_hour_occupancy_rate (mu_{s,h}). Empirical occupancy probability for station s at hour h, computed exclusively on the training split (1,320 lookup table)."),
        ("4. Session Dwell Prior (1)", "station_avg_session_duration. Historical average connection duration per station (80-400 min), distinguishing rapid turnover vs long-dwell chargers."),
        ("Information Provided to Models", "Combined, these 8 features provide the models with calendar time, immediate state inertia, station-specific baselines, and dwell time context.")
    ]
    add_two_column_slide(
        prs, "Methodology: 8-Feature Engineering Hierarchy", f_bullets,
        "artifacts/progress_figures/fig4_feature_hierarchy.png",
        "We engineered 8 features across four logical categories: Temporal features capture human commute schedules; Historical lags capture session inertia—since EVs stay plugged in for hours; Station priors capture charger-specific popularity; and Session dwell captures average parking duration."
    )

    # =========================================================================
    # SLIDE 10: Methodology 3/3 (Classification Setup & Imbalance)
    # =========================================================================
    print("Building Slide 10: Methodology - Classification Setup...")
    setup_bullets = [
        ("Binary Classification Formulation", "Target variable: target_occupancy in {0, 1}. Class 0 = Available, Class 1 = Occupied."),
        ("Severe Class Imbalance", "Training set (1,662,925 intervals): 1,604,113 Available (96.46%) vs 58,812 Occupied (3.54%). Occupied instances are rare!"),
        ("Why Accuracy is Misleading", "A naive model predicting 'Available' 100% of the time achieves 94.4% to 96.5% accuracy but is completely useless (F1 = 0.0000)."),
        ("Primary Evaluation Metrics", "Precision-Recall AUC (PR-AUC) and F1-Score are prioritized because they directly measure minority class detection without negative-class distortion."),
        ("Chronological Partitions", "Train: Oct 2019 - Jun 2021 (89.34%) | Val: Jul 2021 - Aug 2021 (6.38%) | Test: Aug 2021 - Sep 2021 (4.28%).")
    ]
    add_two_column_slide(
        prs, "Methodology: Classification Setup & Imbalance", setup_bullets,
        "artifacts/progress_figures/fig3_class_splits.png",
        "In our classification setup, only 3.54% of training intervals are occupied. Under such extreme imbalance, standard accuracy is misleading because predicting available every time yields 96% accuracy with zero predictive utility. Therefore, we emphasize PR-AUC and F1-score as our primary evaluation metrics."
    )

    # =========================================================================
    # SLIDE 11: Design / Workflow (Research Progression)
    # =========================================================================
    print("Building Slide 11: Workflow...")
    wf_paragraphs = [
        ("STEP 1: DATA PREPARATION & CHRONOLOGICAL PARTITIONING",
         "Discretized 1.86M intervals from Caltech ACN sessions, built 30-min targets, and established leak-free chronological Train, Validation, and Test splits."),
        ("STEP 2: MODEL 1 — LOGISTIC REGRESSION (LINEAR BASELINE)",
         "Trained Logistic Regression with balanced class weighting to establish an interpretable linear baseline and reference performance benchmark."),
        ("STEP 3: ANALYSIS & TRANSITION RATIONALE",
         "Analyzed Model 1: achieved high recall (79.3%) but moderate precision (62.5%), showing that linear boundaries trigger false alarms and miss non-linear interactions."),
        ("STEP 4: MODEL 2 — RANDOM FOREST (NONLINEAR ENSEMBLE)",
         "Trained Random Forest (100 trees) to capture non-linear thresholds and feature interactions between time, station identity, and rolling occupancy."),
        ("STEP 5: ONGOING & PLANNED EXPERIMENTATION",
         "Current stage: Random Forest completed. Next stage: evaluate additional advanced tree and boosting approaches before final model selection.")
    ]
    add_full_width_slide(
        prs, "Design / Workflow: Iterative Research Progression", wf_paragraphs,
        "This workflow diagram illustrates our iterative research approach: we started with a simple linear baseline, analyzed its limitations, moved to a tree ensemble to capture non-linear interactions, and are continuing with further planned model experiments."
    )

    # =========================================================================
    # SLIDE 12: Model 1 — Logistic Regression (Results Slide 1/3)
    # =========================================================================
    print("Building Slide 12: Model 1 - Logistic Regression...")
    m1_bullets = [
        ("Why Selected First?", "Logistic Regression was chosen as our initial model because:"),
        ("  * Simplicity & Interpretability", "Transparent linear model with well-understood probability formulation."),
        ("  * Computational Efficiency", "Trained in just 2.10 seconds, providing rapid iterative feedback."),
        ("  * Reference Benchmark", "Establishes a solid baseline to determine how well linear boundaries perform before attempting complex architectures."),
        ("Model Setup & Class Weighting", "StandardScaler normalization + class_weight='balanced' with lbfgs solver."),
        ("Actual Test Metrics (Exp D)", ""),
        ("  * PR-AUC: 0.7451 | F1-Score: 0.6990", "Primary discrimination and harmonic balance."),
        ("  * Precision: 0.6248 | Recall: 0.7932", "High recall (catches ~79% of occupied states), but lower precision causes false alarms."),
        ("  * ROC-AUC: 0.9596 | Accuracy: 0.9617", "High overall accuracy and area under ROC curve."),
        ("Key Interpretation", "Logistic Regression provides a respectable baseline, but its linear formulation produces a noticeable number of false positives (2,128 FP on test).")
    ]
    add_two_column_slide(
        prs, "Model Development: Model 1 - Logistic Regression", m1_bullets,
        "artifacts/progress_figures/fig7_pr_curves_logreg_rf.png",
        "I started with Logistic Regression because I wanted a simple and interpretable baseline. It gives me a reference point before trying models that can capture nonlinear relationships. It achieved a test PR-AUC of 0.7451 and F1 of 0.6990. While its recall is strong at 79.3%, its precision is 62.5%, indicating that a linear boundary flags too many false alarms."
    )

    # =========================================================================
    # SLIDE 13: Model 2 - Random Forest (Results Slide 2/3)
    # =========================================================================
    print("Building Slide 13: Model 2 - Random Forest...")
    m2_bullets = [
        ("Why Move to Random Forest?", "Based on the characteristics of the EV occupancy problem:"),
        ("  * Non-Linear Relationships", "Charging behavior involves sharp time-of-day thresholds and dwell limits that linear boundaries cannot cleanly separate."),
        ("  * Feature Interactions", "Crucial interactions exist between station identity, time of day, and immediate occupancy inertia."),
        ("  * Ensemble Decision Trees", "Random Forest builds 100 decorrelated trees, reducing variance and modeling complex multi-dimensional partitions."),
        ("Model Setup", "n_estimators=100, max_depth=12, min_samples_leaf=50, class_weight='balanced'."),
        ("Actual Test Metrics (Exp D)", ""),
        ("  * PR-AUC: 0.7860 (+0.0409 gain!)", "Substantial improvement in precision-recall discrimination."),
        ("  * F1-Score: 0.7656 (+0.0666 gain!)", "Marked increase in classification quality."),
        ("  * Precision: 0.7652 (+0.1404 gain!)", "False positives reduced by more than half (from 2,128 down to 1,050!)."),
        ("  * Recall: 0.7661 | ROC-AUC: 0.9766", "Maintains strong recall with balanced precision.")
    ]
    add_two_column_slide(
        prs, "Model Development: Model 2 - Random Forest", m2_bullets,
        "artifacts/progress_figures/fig6_model_comparison.png",
        "After establishing the Logistic Regression result, I moved to Random Forest because charging occupancy may depend on nonlinear interactions between time and historical occupancy features. Random Forest allows me to test whether a more flexible model improves the result. On the test set, Random Forest achieved PR-AUC of 0.7860 and F1 of 0.7656, reducing false alarms by more than half."
    )

    # =========================================================================
    # SLIDE 14: Model Comparison & Discussion (Results Slide 3/3)
    # =========================================================================
    print("Building Slide 14: Model Comparison...")
    comp_bullets = [
        ("Quantitative Test Comparison (Exp D Full Features, N = 79,585)", ""),
        ("  * Metric: PR-AUC (Primary)", "Logistic Reg: 0.7451  |  Random Forest: 0.7860  (Delta: +0.0409)"),
        ("  * Metric: F1-Score", "Logistic Reg: 0.6990  |  Random Forest: 0.7656  (Delta: +0.0666)"),
        ("  * Metric: Precision", "Logistic Reg: 0.6248  |  Random Forest: 0.7652  (Delta: +0.1404)"),
        ("  * Metric: Recall", "Logistic Reg: 0.7932  |  Random Forest: 0.7661  (Delta: -0.0271)"),
        ("  * Metric: ROC-AUC", "Logistic Reg: 0.9596  |  Random Forest: 0.9766  (Delta: +0.0170)"),
        ("  * Metric: Accuracy", "Logistic Reg: 0.9617  |  Random Forest: 0.9737  (Delta: +0.0120)"),
        ("What Did We Learn?", ""),
        ("  1. Non-Linearity Matters", "Tree structures significantly outperform linear hyperplanes in capturing occupancy boundaries."),
        ("  2. False Alarm Reduction", "Random Forest improved Precision by +14.04%, resolving Logistic Regression's high false positive rate."),
        ("  3. Experimentation Ongoing", "Random Forest is our latest evaluated model; final model selection is NOT made yet.")
    ]
    add_two_column_slide(
        prs, "Results & Discussion: Logistic Regression vs. Random Forest", comp_bullets,
        "artifacts/progress_figures/fig6_model_comparison.png",
        "I am not selecting the final model yet. This comparison tells me how the second model behaves relative to the first. Random Forest improved PR-AUC by plus 0.0409 and F1 by plus 0.0666. It cut false alarms substantially. This empirical result proves that non-linear modeling is beneficial, justifying further model exploration."
    )

    # =========================================================================
    # SLIDE 15: Conclusion and Future Scope (Section 8 — 1 slide)
    # =========================================================================
    print("Building Slide 15: Conclusion and Future Scope...")
    concl_items = [
        ("Data Pipeline Completed", "Discretized 1,861,310 intervals from 55 Caltech charging stations and built leak-free 30-minute prediction targets."),
        ("8 Features Engineered", "Designed temporal, historical lag, station-hour priors, and dwell time features strictly respecting chronological boundaries."),
        ("Baseline Established (LogReg)", "Logistic Regression provided initial linear reference performance (PR-AUC 0.7451, F1 0.6990)."),
        ("Non-Linear Gain Confirmed (RF)", "Random Forest improved PR-AUC to 0.7860 and F1 to 0.7656, confirming the value of non-linear tree ensembles."),
        ("No Final Model Selected Yet", "Work is ongoing; final model selection is intentionally deferred pending further experimentation.")
    ]
    future_items = [
        ("Evaluate Additional Suitable ML Models", "Investigate whether additional advanced tree-based or boosting architectures provide further gains."),
        ("Unified Comparative Framework", "Benchmark all candidate models under identical chronological validation partitions."),
        ("Generalization & Error Analysis", "Analyze edge cases, unexpected session disconnects, and transition prediction accuracy."),
        ("Final Held-Out Test Evaluation", "Execute unbiased final evaluation on held-out test data for the end-semester presentation."),
        ("Final Champion Model Selection", "Select the definitive production model based strictly on comparative empirical evidence.")
    ]
    add_two_column_cards_slide(
        prs, "Conclusion and Future scope",
        "Summary of Progress Completed", concl_items,
        "Next Steps Before End-Semester", future_items,
        "In conclusion, we have completed data preprocessing, feature engineering, and trained two distinct model families. Random Forest confirmed that non-linear decision trees capture occupancy patterns better than a linear baseline. Before the final presentation, we will evaluate additional suitable boosting-based and advanced tree models to determine the champion architecture."
    )

    # =========================================================================
    # SLIDE 16: References / Bibliography (Section 9 — 1 slide)
    # =========================================================================
    print("Building Slide 16: References...")
    ref_items = [
        "1.  Z. J. Lee, G. X. Li, and S. H. Low, 'ACN-Data: Analysis and Applications of an Open EV Charging Dataset,' in Proc. 10th ACM Int. Conf. on Future Energy Systems (e-Energy), Phoenix, AZ, USA, 2019, pp. 139–149.",
        "2.  L. Breiman, 'Random Forests,' Machine Learning, vol. 45, no. 1, pp. 5–32, 2001.",
        "3.  D. R. Cox, 'The Regression Analysis of Binary Sequences,' Journal of the Royal Statistical Society: Series B (Methodological), vol. 20, no. 2, pp. 215–232, 1958.",
        "4.  T. Saito and M. Rehmsmeier, 'The Precision-Recall Plot Is More Informative than the ROC Plot When Evaluating Binary Classifiers on Imbalanced Datasets,' PLOS ONE, vol. 10, no. 3, p. e0118432, 2015.",
        "5.  F. Pedregosa et al., 'Scikit-learn: Machine Learning in Python,' Journal of Machine Learning Research, vol. 12, pp. 2825–2830, 2011.",
        "6.  S. Deb, K. Tammi, K. Kalita, and P. Mahanta, 'Impact of Electric Vehicle Charging Station on Battery Life and Grid: A Review,' IEEE Access, vol. 6, pp. 6412–6425, 2018."
    ]
    add_full_width_slide(
        prs, "References / Bibliography", ref_items,
        "These are the foundational academic literature references for our ongoing research, covering the ACN dataset, Random Forest, Logistic Regression, Scikit-learn, and imbalanced classification metrics."
    )

    # =========================================================================
    # SLIDE 17: Thank You! / Q&A
    # =========================================================================
    print("Building Slide 17: Thank You...")
    thank_slide = prs.slides.add_slide(prs.slide_layouts[1])
    
    title = thank_slide.shapes.title
    title.text = "Thank You!"
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = COLOR_NAVY
    title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    body = thank_slide.placeholders[1]
    body.left = Inches(1.5)
    body.top = Inches(2.2)
    body.width = Inches(10.33)
    body.height = Inches(4.2)
    
    tf = body.text_frame
    tf.word_wrap = True
    tf.clear()
    
    p0 = tf.paragraphs[0]
    p0.text = "Questions & Interim Feedback"
    p0.font.size = Pt(28)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_BLUE
    p0.alignment = PP_ALIGN.CENTER
    p0.space_after = Pt(24)
    
    p1 = tf.add_paragraph()
    p1.text = "EVConnect: Machine Learning-Based Prediction of EV Charging Station Occupancy"
    p1.font.size = Pt(18)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_NAVY
    p1.alignment = PP_ALIGN.CENTER
    p1.space_after = Pt(10)
    
    p2 = tf.add_paragraph()
    p2.text = "Presented by: Kunal Singh  |  Roll No: [Insert Roll No / 23223xxx]\nMCA - V Semester  |  Department of Computer Applications\nNational Institute of Technology (NIT), Raipur"
    p2.font.size = Pt(16)
    p2.font.color.rgb = COLOR_DARK
    p2.alignment = PP_ALIGN.CENTER
    p2.space_after = Pt(16)
    
    p3 = tf.add_paragraph()
    p3.text = "Evaluated Models: Logistic Regression  |  Random Forest  |  Further Models in Progress"
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLOR_MUTED
    p3.alignment = PP_ALIGN.CENTER
    
    thank_slide.notes_slide.notes_text_frame.text = (
        "Thank you respected committee members and supervisor. I welcome your questions, suggestions, "
        "and feedback on our interim progress and planned next steps."
    )

    # Save presentation
    print(f"Saving presentation to {DESKTOP_OUT}...")
    prs.save(DESKTOP_OUT)
    
    print(f"Saving presentation to {WORKSPACE_OUT}...")
    prs.save(WORKSPACE_OUT)
    
    print(f"Total slides generated: {len(prs.slides)}")
    print("SUCCESS: Progress Presentation built cleanly!")

if __name__ == "__main__":
    build_presentation()
